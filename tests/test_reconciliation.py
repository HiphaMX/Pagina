import pytest
from datetime import datetime
from fastapi.testclient import TestClient

from app.main import app
from app.api.deps import get_current_active_user
from app.schemas.user import User as UserSchema
from app.core.database import SessionLocal, Base, engine
from app.models.reconciliation import CFDIInvoice, BankTransaction
from app.services.reconciliation.cfdi_parser import parse_cfdi_xml
from app.services.reconciliation.bbva_parser import parse_bbva_statement
from app.services.reconciliation.matching_engine import conciliar_movimientos
from app.services.reconciliation.excel_exporter import export_reconciliation_excel
from app.services.reconciliation.pdf_reporter import generate_discrepancies_pdf
from app.services.reconciliation.backfill import import_historical_excel

client = TestClient(app)

mock_admin = UserSchema(
    id=1,
    email="hola@hipha.mx",
    is_active=True,
    is_superuser=True,
    full_name="Administrador Hipha"
)

SAMPLE_CFDI_XML = b"""<?xml version="1.0" encoding="utf-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4" xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
    Fecha="2026-05-10T12:00:00" Total="1160.00" SubTotal="1000.00" MetodoPago="PUE" FormaPago="03" TipoDeComprobante="I">
    <cfdi:Emisor Rfc="PROV123456ABC" Nombre="PROVEEDORA DIGITAL SA DE CV" />
    <cfdi:Receptor Rfc="HMA1803098A3" Nombre="HIPHA MARKETING DIGITAL" />
    <cfdi:Conceptos>
        <cfdi:Concepto ClaveProdServ="81111500" Cantidad="1" Descripcion="Licencia Cloud Hosting" Importe="1000.00" />
    </cfdi:Conceptos>
    <cfdi:Impuestos TotalImpuestosTrasladados="160.00">
        <cfdi:Traslados>
            <cfdi:Traslado Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000" Importe="160.00" />
        </cfdi:Traslados>
    </cfdi:Impuestos>
    <cfdi:Complemento>
        <tfd:TimbreFiscalDigital UUID="AAAAAAAA-BBBB-CCCC-DDDD-EEEEEEEEEEEE" FechaTimbrado="2026-05-10T12:05:00" />
    </cfdi:Complemento>
</cfdi:Comprobante>
"""

SAMPLE_BBVA_CSV = b"""Fecha,Concepto,Cargo,Abono,Saldo
12/05/2026,SPEI PAGO PROVEEDORA DIGITAL REF 49201,1160.00,,45000.00
14/05/2026,COMISION POR TRANSFERENCIA INTERBANCARIA,5.80,,44994.20
15/05/2026,DEPOSITO EN EFECTIVO,,5000.00,49994.20
"""

@pytest.fixture(autouse=True)
def setup_db_and_auth():
    Base.metadata.create_all(bind=engine)
    app.dependency_overrides[get_current_active_user] = lambda: mock_admin
    yield
    app.dependency_overrides.pop(get_current_active_user, None)


def test_cfdi_parser():
    # Probar como EGRESO (Receptor es Hipha)
    parsed = parse_cfdi_xml(SAMPLE_CFDI_XML, mi_rfc="HMA1803098A3")
    assert parsed["uuid"] == "AAAAAAAA-BBBB-CCCC-DDDD-EEEEEEEEEEEE"
    assert parsed["tipo"] == "EGRESO"
    assert parsed["total"] == 1160.00
    assert parsed["subtotal"] == 1000.00
    assert parsed["iva_trasladado"] == 160.00
    assert parsed["rfc_emisor"] == "PROV123456ABC"
    assert "PROVEEDORA DIGITAL" in parsed["nombre_emisor"]

    # Probar como INGRESO si el emisor fuera Hipha
    parsed_ingreso = parse_cfdi_xml(SAMPLE_CFDI_XML, mi_rfc="PROV123456ABC")
    assert parsed_ingreso["tipo"] == "INGRESO"


def test_bbva_parser():
    txs = parse_bbva_statement(SAMPLE_BBVA_CSV, "movimientos_mayo.csv")
    assert len(txs) == 3

    # Primera tx: Cargo 1160
    assert txs[0]["monto"] == 1160.00
    assert txs[0]["tipo"] == "EGRESO"
    assert "PROVEEDORA DIGITAL" in txs[0]["concepto"]
    assert txs[0]["periodo_mes"] == "2026-05"
    assert len(txs[0]["id_transaccion"]) == 64  # SHA-256

    # Segunda tx: Comisión 5.80
    assert txs[1]["monto"] == 5.80
    assert txs[1]["tipo"] == "EGRESO"

    # Tercera tx: Abono 5000
    assert txs[2]["monto"] == 5000.00
    assert txs[2]["tipo"] == "INGRESO"


def test_bbva_parser_spanish_months_and_agosto():
    bbva_agosto_csv = b"""Fecha,Concepto / Referencia,Cargo,Abono,Saldo
01/AGO/2026,PAGO PROVEEDOR SERVICIOS CLOUD,2500.00,,42500.00
15-AGO-26,DEPOSITO CLIENTE HONORARIOS,,15000.00,57500.00
28/Ago/2026,COMISION MANEJO DE CUENTA,150.00,,57350.00
"""
    txs = parse_bbva_statement(bbva_agosto_csv, "movimientos_agosto.csv")
    assert len(txs) == 3

    assert txs[0]["periodo_mes"] == "2026-08"
    assert txs[0]["monto"] == 2500.00
    assert txs[0]["tipo"] == "EGRESO"

    assert txs[1]["periodo_mes"] == "2026-08"
    assert txs[1]["monto"] == 15000.00
    assert txs[1]["tipo"] == "INGRESO"

    assert txs[2]["periodo_mes"] == "2026-08"
    assert txs[2]["monto"] == 150.00
    assert txs[2]["tipo"] == "EGRESO"


def test_bbva_parser_semicolon_and_preamble():
    bbva_semi_csv = """BBVA Bancomer S.A.
Cuenta: 0123456789
Periodo consultado: 01/08/2026 al 31/08/2026

Fecha de operacion;Concepto;Retiro;Deposito;Saldo
05/08/2026;COMPRA SUPERMERCADO;(1.250,50);;38749,50
10/08/2026;TRANSFERENCIA SPEI ENTRADA;;5.000,00;43749,50
TOTAL DE MOVIMIENTOS: 2
""".encode("latin-1")

    txs = parse_bbva_statement(bbva_semi_csv, "bbva_export_netcash.csv")
    assert len(txs) == 2

    # Retiro con parentesis y formato decimal europeo
    assert txs[0]["monto"] == 1250.50
    assert txs[0]["tipo"] == "EGRESO"
    assert txs[0]["periodo_mes"] == "2026-08"

    # Deposito
    assert txs[1]["monto"] == 5000.00
    assert txs[1]["tipo"] == "INGRESO"
    assert txs[1]["periodo_mes"] == "2026-08"


def test_matching_engine():
    parsed_cfdi = parse_cfdi_xml(SAMPLE_CFDI_XML, mi_rfc="HMA1803098A3")
    txs = parse_bbva_statement(SAMPLE_BBVA_CSV, "movimientos_mayo.csv")

    facturas_pool = [parsed_cfdi]
    reconciled = conciliar_movimientos(txs, facturas_pool)

    assert len(reconciled) == 3

    # El cargo de 1160 debe ser Nivel 1 (CONCILIADO 100%) porque el monto coincide (1160),
    # fecha está en ventana (10 a 12 mayo = +2 días) y el nombre 'PROVEEDORA DIGITAL' está en el concepto bancario.
    tx_1160 = next(t for t in reconciled if t["monto"] == 1160.00)
    assert tx_1160["status_conciliacion"] == "CONCILIADO"
    assert tx_1160["uuid_cfdi"] == "AAAAAAAA-BBBB-CCCC-DDDD-EEEEEEEEEEEE"
    assert tx_1160["confianza_score"] == 1.0

    # La comisión de 5.80 debe ser SIN_CFDI
    tx_comision = next(t for t in reconciled if t["monto"] == 5.80)
    assert tx_comision["status_conciliacion"] == "SIN_CFDI"
    assert "comisión" in tx_comision["nota_revision"].lower() or "bancario" in tx_comision["nota_revision"].lower()


def test_exporters():
    txs = parse_bbva_statement(SAMPLE_BBVA_CSV, "movimientos_mayo.csv")
    parsed_cfdi = parse_cfdi_xml(SAMPLE_CFDI_XML, mi_rfc="HMA1803098A3")
    reconciled = conciliar_movimientos(txs, [parsed_cfdi])

    # Excel
    xlsx_bytes = export_reconciliation_excel(reconciled, mes_label="2026-05")
    assert len(xlsx_bytes) > 1000

    # PDF Faltantes
    pdf_bytes = generate_discrepancies_pdf(reconciled, mes_label="2026-05")
    assert len(pdf_bytes) > 500
    assert pdf_bytes.startswith(b"%PDF")


def test_api_reconciliation_endpoints():
    # 1. Overview
    res = client.get("/api/dashboard/reconciliation/overview?account_rfc=DEGF851127TK1")
    assert res.status_code == 200
    data = res.json()
    assert "kpis" in data
    assert "selected_month" in data
    assert data["entity_name"] == "HIPHA"

    # AMDI Overview
    res_amdi = client.get("/api/dashboard/reconciliation/overview?account_rfc=MEHA850118Q96")
    assert res_amdi.status_code == 200
    assert res_amdi.json()["entity_name"] == "AMDI"

    # 2. Transactions
    res_tx = client.get("/api/dashboard/reconciliation/transactions?account_rfc=DEGF851127TK1")
    assert res_tx.status_code == 200
    assert isinstance(res_tx.json(), list)

    # 3. Pending invoices
    res_inv = client.get("/api/dashboard/reconciliation/pending-invoices?account_rfc=DEGF851127TK1")
    assert res_inv.status_code == 200
    assert isinstance(res_inv.json(), list)


def test_multi_entity_isolation():
    db = SessionLocal()
    try:
        # Create one transaction for HIPHA and one for AMDI
        tx_hipha = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx_test_hipha_001",
            fecha=datetime(2026, 5, 20),
            periodo_mes="2026-05",
            concepto="PAGO SERVICIO HIPHA DOMINIO",
            monto=500.0,
            tipo="EGRESO",
            saldo=10000.0,
            status_conciliacion="SIN_CFDI",
            area_proyecto="HIPHA"
        )
        tx_amdi = BankTransaction(
            account_rfc="MEHA850118Q96",
            id_transaccion="tx_test_amdi_001",
            fecha=datetime(2026, 5, 21),
            periodo_mes="2026-05",
            concepto="COMPRA INSUMOS AMDI FABRICA",
            monto=1200.0,
            tipo="EGRESO",
            saldo=25000.0,
            status_conciliacion="SIN_CFDI",
            area_proyecto="AMDI"
        )
        db.add(tx_hipha)
        db.add(tx_amdi)
        db.commit()

        # Query HIPHA transactions
        res_h = client.get("/api/dashboard/reconciliation/transactions?account_rfc=DEGF851127TK1&month=2026-05")
        assert res_h.status_code == 200
        txs_h = res_h.json()
        assert any(t["id_transaccion"] == "tx_test_hipha_001" for t in txs_h)
        assert not any(t["id_transaccion"] == "tx_test_amdi_001" for t in txs_h)

        # Query AMDI transactions
        res_a = client.get("/api/dashboard/reconciliation/transactions?account_rfc=MEHA850118Q96&month=2026-05")
        assert res_a.status_code == 200
        txs_a = res_a.json()
        assert any(t["id_transaccion"] == "tx_test_amdi_001" for t in txs_a)
        assert not any(t["id_transaccion"] == "tx_test_hipha_001" for t in txs_a)

    finally:
        db.query(BankTransaction).filter(BankTransaction.id_transaccion.in_(["tx_test_hipha_001", "tx_test_amdi_001"])).delete(synchronize_session=False)
        db.commit()
        db.close()


def test_14_column_excel_export_structure():
    import io
    import openpyxl

    sample_txs = [
        {
            "fecha": "2026-05-15",
            "area_proyecto": "AMDI Producción",
            "uuid_cfdi": "11111111-2222-3333-4444-555555555555",
            "contraparte": "ACEROS DE MEXICO SA",
            "concepto": "PAGO LAMINA GALVANIZADA",
            "metodo_pago": "PUE",
            "monto": 1160.0,
            "tipo": "EGRESO",
            "status_conciliacion": "CONCILIADO",
            "subtotal": 1000.0,
            "iva_egresos": 160.0,
            "retencion_isr": 12.5,
            "tipo_categoria": "GASTO"
        }
    ]

    xlsx_bytes = export_reconciliation_excel(sample_txs, mes_label="2026-05", account_rfc="MEHA850118Q96")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # Check Title
    assert "AMDI" in ws["A1"].value

    # Check 14 Headers on row 3
    expected_headers = [
        "FECHA",
        "ÁREA / PROYECTO",
        "FOLIO FISCAL",
        "CLIENTE / PROVEEDOR",
        "CONCEPTO",
        "MÉTODO DE PAGO",
        "INGRESOS",
        "EGRESOS",
        "CFDI",
        "TIPO",
        "SUBTOTAL",
        "IVA INGRESOS",
        "RETENCIÓN ISR",
        "IVA EGRESOS"
    ]

    actual_headers = [ws.cell(row=3, column=c).value for c in range(1, 15)]
    assert actual_headers == expected_headers

    # Check Data on row 4
    assert ws.cell(row=4, column=1).value == "2026-05-15"
    assert ws.cell(row=4, column=2).value == "AMDI Producción"
    assert ws.cell(row=4, column=3).value == "11111111-2222-3333-4444-555555555555"
    assert ws.cell(row=4, column=4).value == "ACEROS DE MEXICO SA"
    assert ws.cell(row=4, column=8).value == 1160.0  # EGRESOS
    assert ws.cell(row=4, column=11).value == 1000.0 # SUBTOTAL
    assert ws.cell(row=4, column=13).value == 12.5   # RETENCIÓN ISR
    assert ws.cell(row=4, column=14).value == 160.0  # IVA EGRESOS

    # Check Formulas on totals row (row 5)
    assert ws.cell(row=5, column=5).value == "TOTALES"
    assert "=SUM(" in str(ws.cell(row=5, column=7).value)
    assert "=SUM(" in str(ws.cell(row=5, column=8).value)
    assert "=SUM(" in str(ws.cell(row=5, column=11).value)
    assert "=SUM(" in str(ws.cell(row=5, column=13).value)
    assert "=SUM(" in str(ws.cell(row=5, column=14).value)


def test_cfdi_retention_isr_iva_separation():
    xml_with_retentions = b"""<?xml version="1.0" encoding="utf-8"?>
    <cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4" xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
        Fecha="2026-05-18T10:00:00" Total="1025.00" SubTotal="1000.00" MetodoPago="PUE" FormaPago="03" TipoDeComprobante="I">
        <cfdi:Emisor Rfc="PROF123456XYZ" Nombre="SERVICIOS PROFESIONALES SC" />
        <cfdi:Receptor Rfc="MEHA850118Q96" Nombre="AMDI DECORACION" />
        <cfdi:Impuestos TotalImpuestosTrasladados="160.00" TotalImpuestosRetenidos="135.00">
            <cfdi:Traslados>
                <cfdi:Traslado Impuesto="002" TipoFactor="Tasa" TasaOCuota="0.160000" Importe="160.00" />
            </cfdi:Traslados>
            <cfdi:Retenciones>
                <cfdi:Retencion Impuesto="001" Importe="100.00" />
                <cfdi:Retencion Impuesto="002" Importe="35.00" />
            </cfdi:Retenciones>
        </cfdi:Impuestos>
        <cfdi:Complemento>
            <tfd:TimbreFiscalDigital UUID="99999999-8888-7777-6666-555555555555" />
        </cfdi:Complemento>
    </cfdi:Comprobante>
    """
    parsed = parse_cfdi_xml(xml_with_retentions, mi_rfc="MEHA850118Q96")
    assert parsed["uuid"] == "99999999-8888-7777-6666-555555555555"
    assert parsed["subtotal"] == 1000.0
    assert parsed["iva_trasladado"] == 160.0
    assert parsed["retencion_isr"] == 100.0
    assert parsed["retencion_iva"] == 35.0
    assert parsed["retenciones"] == 135.0


def test_historical_backfill_csv_14_columns():
    db = SessionLocal()
    sample_csv = (
        b"FECHA,\xc1REA / PROYECTO,FOLIO FISCAL,CLIENTE / PROVEEDOR,CONCEPTO,M\xc9TODO DE PAGO,INGRESOS,EGRESOS,CFDI,TIPO,SUBTOTAL,IVA INGRESOS,RETENCI\xd3N ISR,IVA EGRESOS\n"
        b"15/01/2026,HIPHA MKT,12345678-1234-1234-1234-123456789012,CLIENTE ABC,CAMPANA DIGITAL GOOGLE,PUE,11600.00,,I,INGRESO,10000.00,1600.00,0.00,0.00\n"
        b"18/01/2026,HIPHA MKT,,BBVA COMISIONES,COMISION MENSUAL CUENTA,PUE,,232.00,,GASTO,200.00,0.00,0.00,32.00\n"
    )

    try:
        res = import_historical_excel(
            sample_csv,
            "ENero.csv",
            db,
            mi_rfc="DEGF851127TK1",
            account_rfc="DEGF851127TK1"
        )
        assert res["success"] is True
        assert res["imported_transactions"] == 2
        assert res["imported_invoices"] == 1

        # Check that the invoice was registered as CONCILIADO
        inv = db.query(CFDIInvoice).filter(CFDIInvoice.uuid == "12345678-1234-1234-1234-123456789012").first()
        assert inv is not None
        assert inv.conciliado is True
        assert inv.total == 11600.00

        # Check transactions
        txs = db.query(BankTransaction).filter(BankTransaction.account_rfc == "DEGF851127TK1", BankTransaction.periodo_mes == "2026-01").all()
        assert len(txs) == 2
        conciliated_tx = next(t for t in txs if t.uuid_cfdi == "12345678-1234-1234-1234-123456789012")
        assert conciliated_tx.status_conciliacion == "CONCILIADO"
        assert conciliated_tx.monto == 11600.00
    finally:
        db.query(CFDIInvoice).filter(CFDIInvoice.uuid == "12345678-1234-1234-1234-123456789012").delete()
        db.query(BankTransaction).filter(BankTransaction.account_rfc == "DEGF851127TK1", BankTransaction.periodo_mes == "2026-01").delete()
        db.commit()
        db.close()


def test_delete_transaction():
    db = SessionLocal()
    try:
        db.query(BankTransaction).filter(BankTransaction.id_transaccion == "tx_test_del_001").delete()
        db.query(CFDIInvoice).filter(CFDIInvoice.uuid == "TEST-DEL-UUID-001").delete()
        db.commit()

        inv = CFDIInvoice(
            uuid="TEST-DEL-UUID-001",
            account_rfc="DEGF851127TK1",
            tipo="EGRESO",
            rfc_emisor="PROV123",
            rfc_receptor="DEGF851127TK1",
            fecha_emision=datetime(2026, 7, 10),
            total=500.0,
            conciliado=True
        )
        db.add(inv)
        db.commit()

        tx = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx_test_del_001",
            fecha=datetime(2026, 7, 10),
            periodo_mes="2026-07",
            concepto="CARGO DE PRUEBA A ELIMINAR",
            monto=500.0,
            tipo="EGRESO",
            uuid_cfdi="TEST-DEL-UUID-001",
            status_conciliacion="CONCILIADO"
        )
        db.add(tx)
        db.commit()
        tx_id = tx.id

        res = client.delete(f"/api/dashboard/reconciliation/transactions/{tx_id}")
        assert res.status_code == 200
        assert res.json()["success"] is True

        # Verificar que la transacción fue borrada
        deleted = db.query(BankTransaction).filter(BankTransaction.id == tx_id).first()
        assert deleted is None

        # Verificar que la factura quedó liberada
        inv_updated = db.query(CFDIInvoice).filter(CFDIInvoice.uuid == "TEST-DEL-UUID-001").first()
        assert inv_updated is not None
        assert inv_updated.conciliado is False
    finally:
        db.query(CFDIInvoice).filter(CFDIInvoice.uuid == "TEST-DEL-UUID-001").delete()
        db.commit()
        db.close()


def test_deduplicate_transactions():
    db = SessionLocal()
    try:
        db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07",
            BankTransaction.monto == 22605.75
        ).delete()
        db.commit()

        tx1 = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx_dup_001",
            fecha=datetime(2026, 7, 22),
            periodo_mes="2026-07",
            concepto="ABONO",
            monto=22605.75,
            tipo="INGRESO",
            area_proyecto="DAM",
            status_conciliacion="SIN_CFDI"
        )
        tx2 = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx_dup_002",
            fecha=datetime(2026, 7, 22),
            periodo_mes="2026-07",
            concepto="ABONO",
            monto=22605.75,
            tipo="INGRESO",
            area_proyecto="DAM",
            status_conciliacion="SIN_CFDI"
        )
        db.add(tx1)
        db.add(tx2)
        db.commit()

        # Debe haber 2 duplicados antes
        dups_before = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07",
            BankTransaction.monto == 22605.75
        ).all()
        assert len(dups_before) == 2

        res = client.post("/api/dashboard/reconciliation/deduplicate?account_rfc=DEGF851127TK1&month=2026-07")
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["deleted_count"] >= 1

        # Debe quedar solo 1
        dups_after = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07",
            BankTransaction.monto == 22605.75
        ).all()
        assert len(dups_after) == 1
    finally:
        db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07",
            BankTransaction.monto == 22605.75
        ).delete()
        db.commit()
        db.close()


def test_cron_sync_mailbox(monkeypatch):
    calls = []

    def mock_sync(db, account_rfc, lookback_days=60):
        calls.append((account_rfc, lookback_days))
        return {
            "success": True,
            "connected": True,
            "message": f"Sincronizado {account_rfc}",
            "processed_xmls": 1,
            "processed_statements": 0
        }

    monkeypatch.setattr("app.api.dashboard.reconciliation_routes.sync_mailbox_invoices", mock_sync)

    res = client.get("/api/dashboard/reconciliation/cron-sync")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "hipha" in data
    assert "amdi" in data
    assert len(calls) == 2
    assert calls[0][0] == "DEGF851127TK1"
    assert calls[1][0] == "MEHA850118Q96"


def test_backfill_agosto_declared_month_and_concept_synthesis():
    with open("tests/fixtures/agosto.csv", "rb") as f:
        content = f.read()

    db = SessionLocal()
    try:
        res = import_historical_excel(content, "agosto.csv", db, mi_rfc="DEGF851127TK1")
        assert res["success"] is True
        assert res["imported_transactions"] == 3
        assert res["transacciones_importadas"] == 3
        assert "2026-07" in res["meses_afectados"]

        txs = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07"
        ).all()
        assert len(txs) >= 3

        # Conceptos no deben estar vacíos
        for t in txs:
            assert t.concepto and len(t.concepto) > 3
            assert t.periodo_mes == "2026-07"

        # Verificar montos de las 3 transacciones importadas
        montos_subset = sorted([float(t.monto) for t in txs if float(t.monto) in [383.88, 4361.60, 22605.75]])
        assert 383.88 in montos_subset
        assert 4361.60 in montos_subset
        assert 22605.75 in montos_subset

        # Verificar endpoint de upload con archivo agosto.csv
        files = [("files", ("agosto.csv", content, "text/csv"))]
        upload_res = client.post("/api/dashboard/reconciliation/upload-files", files=files, data={"account_rfc": "DEGF851127TK1"})
        assert upload_res.status_code == 200
        data_up = upload_res.json()
        assert data_up["success"] is True
        assert data_up["processed_statements"] == 3
        assert "2026-07" in data_up["detected_months"]
    finally:
        db.close()


def test_parse_bank_pdf_statement_bbva_comprobante():
    from app.services.reconciliation.pdf_parser import parse_bank_pdf_statement

    with open("tests/fixtures/BBVA.pdf", "rb") as f:
        pdf_bytes = f.read()

    txs = parse_bank_pdf_statement(pdf_bytes, "BBVA.pdf", account_rfc="MEHA850118Q96")
    assert len(txs) == 1
    t = txs[0]
    assert t["monto"] == 2000.00
    assert t["tipo"] == "EGRESO"
    assert t["periodo_mes"] == "2026-10"
    assert "JAIME OMAR RODRIGUEZ ALCALA" in t["concepto"]
    assert "DESINSTALACION LETRERO DAM" in t["concepto"]
    assert t["area_proyecto"] == "DAM"
    assert t["referencia_rastreo"] == "BNET01002610090025425885"


def test_upload_bank_pdf_statement_endpoint():
    with open("tests/fixtures/BBVA.pdf", "rb") as f:
        pdf_bytes = f.read()

    db = SessionLocal()
    try:
        files = [("files", ("BBVA.pdf", pdf_bytes, "application/pdf"))]
        upload_res = client.post(
            "/api/dashboard/reconciliation/upload-files",
            files=files,
            data={"account_rfc": "MEHA850118Q96"}
        )
        assert upload_res.status_code == 200
        data_up = upload_res.json()
        assert data_up["success"] is True
        assert data_up["processed_statements"] >= 1
        assert "2026-10" in data_up["detected_months"]

        tx = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "MEHA850118Q96",
            BankTransaction.periodo_mes == "2026-10",
            BankTransaction.monto == 2000.00
        ).first()
        assert tx is not None
        assert tx.tipo == "EGRESO"
        assert tx.area_proyecto == "DAM"
        assert "JAIME OMAR RODRIGUEZ ALCALA" in tx.concepto
    finally:
        db.close()


def test_parse_bank_pdf_statement_bbva_pyme_julio():
    from app.services.reconciliation.pdf_parser import parse_bank_pdf_statement

    with open("tests/fixtures/Julio_FDG.pdf", "rb") as f:
        pdf_bytes = f.read()

    txs = parse_bank_pdf_statement(pdf_bytes, "Julio FDG.pdf", account_rfc="DEGF851127TK1")
    assert len(txs) == 5

    # Validar periodos y RFC
    for t in txs:
        assert t["periodo_mes"] == "2026-07"
        assert t["account_rfc"] == "DEGF851127TK1"
        assert t["banco"] == "BBVA"

    ingresos = [t for t in txs if t["tipo"] == "INGRESO"]
    egresos = [t for t in txs if t["tipo"] == "EGRESO"]

    assert len(ingresos) == 2
    assert len(egresos) == 3

    assert sum(t["monto"] for t in ingresos) == pytest.approx(26967.35, 0.01)
    assert sum(t["monto"] for t in egresos) == pytest.approx(26967.35, 0.01)

    # Validar detección de proyectos (ej. Axis -> HEALTHYICE)
    axis_tx = next(t for t in ingresos if "Axis" in t["concepto"])
    assert axis_tx["area_proyecto"] == "HEALTHYICE"
    assert axis_tx["monto"] == 4361.60


def test_upload_bank_pdf_julio_endpoint():
    with open("tests/fixtures/Julio_FDG.pdf", "rb") as f:
        pdf_bytes = f.read()

    db = SessionLocal()
    try:
        db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07"
        ).delete()
        db.commit()

        files = [("files", ("Julio FDG.pdf", pdf_bytes, "application/pdf"))]
        upload_res = client.post(
            "/api/dashboard/reconciliation/upload-files",
            files=files,
            data={"account_rfc": "DEGF851127TK1"}
        )
        assert upload_res.status_code == 200
        data_up = upload_res.json()
        assert data_up["success"] is True
        assert data_up["processed_statements"] == 5
        assert "2026-07" in data_up["detected_months"]

        txs = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.periodo_mes == "2026-07"
        ).all()
        assert len(txs) == 5
    finally:
        db.close()


def test_parse_bank_pdf_statement_mercadopago_credit_card():
    from app.services.reconciliation.pdf_parser import parse_bank_pdf_statement

    with open("tests/fixtures/ML_agosto_2026.pdf", "rb") as f:
        pdf_bytes = f.read()

    txs = parse_bank_pdf_statement(pdf_bytes, "ML_agosto_2026.pdf", account_rfc="DEGF851127TK1")
    assert len(txs) == 6

    # Todas las compras con tarjeta son EGRESOS del periodo 2026-08
    for t in txs:
        assert t["tipo"] == "EGRESO"
        assert t["banco"] == "TC MERCADO PAGO"
        assert t["periodo_mes"] == "2026-08"
        assert t["account_rfc"] == "DEGF851127TK1"

    # La suma de las compras debe ser exactamente el total a pagar del periodo: $4,421.29
    total_compras = sum(t["monto"] for t in txs)
    assert total_compras == pytest.approx(4421.29, 0.01)

    # Validar comercios individuales
    conceptos = [t["concepto"] for t in txs]
    assert any("Google One" in c for c in conceptos)
    assert any("GOOGLE CLOUD" in c for c in conceptos)
    assert any("WALMART" in c for c in conceptos)
    assert any("NEUBOX" in c for c in conceptos)
    assert any("MERCADOLIBRE" in c for c in conceptos)
    assert any("MP ECOMMERCE" in c for c in conceptos)

    # Asegurar que saldos anteriores y pagos de la tarjeta no se extraen como transacciones
    assert not any("Saldo al corte" in c for c in conceptos)
    assert not any("Pago del resumen" in c for c in conceptos)


def test_upload_bank_pdf_mercadopago_and_cross_match():
    with open("tests/fixtures/ML_agosto_2026.pdf", "rb") as f:
        pdf_bytes = f.read()

    db = SessionLocal()
    try:
        # 1. Crear facturas CFDI de gastos para cruzar con las compras de la tarjeta
        cfdi_gcloud = CFDIInvoice(
            uuid="CFDI-GCLOUD-TEST-0001",
            account_rfc="DEGF851127TK1",
            area_proyecto="HIPHA",
            tipo="EGRESO",
            rfc_emisor="GCM150101XYZ",
            nombre_emisor="GOOGLE CLOUD MEXICO S DE RL DE CV",
            rfc_receptor="DEGF851127TK1",
            nombre_receptor="FRANCISCO DE JESUS DELGADILLO GARCIA",
            fecha_emision=datetime(2026, 8, 1, 10, 0),
            subtotal=260.12,
            iva_trasladado=41.62,
            retenciones=0.0,
            total=301.74,
            metodo_pago="PUE",
            conceptos_resumen="Consumo Cloud Hosting GCP",
            conciliado=False
        )
        cfdi_walmart = CFDIInvoice(
            uuid="CFDI-WALMART-TEST-0002",
            account_rfc="DEGF851127TK1",
            area_proyecto="HIPHA",
            tipo="EGRESO",
            rfc_emisor="NWM9709244W4",
            nombre_emisor="NUEVA WAL MART DE MEXICO S DE RL DE CV",
            rfc_receptor="DEGF851127TK1",
            nombre_receptor="FRANCISCO DE JESUS DELGADILLO GARCIA",
            fecha_emision=datetime(2026, 8, 4, 18, 30),
            subtotal=1197.41,
            iva_trasladado=191.59,
            retenciones=0.0,
            total=1389.00,
            metodo_pago="PUE",
            conceptos_resumen="Insumos Oficina Walmart",
            conciliado=False
        )
        db.merge(cfdi_gcloud)
        db.merge(cfdi_walmart)
        db.commit()

        # 2. Subir estado de cuenta de Mercado Pago
        files = [("files", ("ML_agosto_2026.pdf", pdf_bytes, "application/pdf"))]
        upload_res = client.post(
            "/api/dashboard/reconciliation/upload-files",
            files=files,
            data={"account_rfc": "DEGF851127TK1"}
        )
        assert upload_res.status_code == 200
        data_up = upload_res.json()
        assert data_up["success"] is True
        assert data_up["processed_statements"] == 6
        assert "2026-08" in data_up["detected_months"]

        # 3. Validar que las transacciones de Google Cloud y Walmart se conciliaron automáticamente
        tx_gcloud = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.monto == 301.74
        ).first()
        assert tx_gcloud is not None
        assert tx_gcloud.status_conciliacion == "CONCILIADO"
        assert tx_gcloud.uuid_cfdi == "CFDI-GCLOUD-TEST-0001"
        assert tx_gcloud.banco == "TC MERCADO PAGO"

        tx_walmart = db.query(BankTransaction).filter(
            BankTransaction.account_rfc == "DEGF851127TK1",
            BankTransaction.monto == 1389.00
        ).first()
        assert tx_walmart is not None
        assert tx_walmart.status_conciliacion == "CONCILIADO"
        assert tx_walmart.uuid_cfdi == "CFDI-WALMART-TEST-0002"
    finally:
        db.close()


def test_pending_invoices_month_filter_and_multi_payment():
    db = SessionLocal()
    try:
        # 1. Crear factura CFDI emitida (INGRESO) de $3,213.20 en agosto 2026
        uuid_tukipa = "CFDI-TUKIPA-MULTI-PAY-001"
        cfdi = CFDIInvoice(
            uuid=uuid_tukipa,
            account_rfc="DEGF851127TK1",
            tipo="INGRESO",
            rfc_emisor="DEGF851127TK1",
            nombre_emisor="FRANCISCO DE JESUS DELGADILLO GARCIA",
            rfc_receptor="MEHA870222DY4",
            nombre_receptor="ABEL ALEJANDRO MEDINA HERNANDEZ",
            fecha_emision=datetime(2026, 8, 28, 10, 0),
            subtotal=2770.00,
            iva_trasladado=443.20,
            retenciones=0.0,
            total=3213.20,
            metodo_pago="PPD",
            conceptos_resumen="Diseño gráfico y publicidad",
            conciliado=False
        )
        db.merge(cfdi)

        # 2. Crear dos transacciones bancarias en BBVA
        tx1 = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx-tukipa-pago1",
            banco="BBVA",
            fecha=datetime(2026, 8, 28),
            concepto="PAGO CUENTA DE TERCERO - Tukipa Marqueting",
            monto=2770.00,
            tipo="EGRESO",  # Clasificado inicialmente como EGRESO por BBVA
            status_conciliacion="SIN_CFDI",
            periodo_mes="2026-08"
        )
        tx2 = BankTransaction(
            account_rfc="DEGF851127TK1",
            id_transaccion="tx-tukipa-pago2",
            banco="BBVA",
            fecha=datetime(2026, 8, 28),
            concepto="PAGO CUENTA DE TERCERO - iva tukipa marketi",
            monto=443.20,
            tipo="EGRESO",
            status_conciliacion="SIN_CFDI",
            periodo_mes="2026-08"
        )
        db.merge(tx1)
        db.merge(tx2)
        db.commit()

        db_tx1 = db.query(BankTransaction).filter(BankTransaction.id_transaccion == "tx-tukipa-pago1").first()
        db_tx2 = db.query(BankTransaction).filter(BankTransaction.id_transaccion == "tx-tukipa-pago2").first()

        # 3. Probar endpoint /pending-invoices con filtro de mes
        res_aug = client.get("/api/dashboard/reconciliation/pending-invoices?month=2026-08&account_rfc=DEGF851127TK1")
        assert res_aug.status_code == 200
        invs_aug = res_aug.json()
        target = next((i for i in invs_aug if i["uuid"] == uuid_tukipa), None)
        assert target is not None
        assert target["total"] == 3213.20
        assert target["tipo"] == "INGRESO"
        assert target["saldo_pendiente"] == 3213.20

        # Verificar que si se filtra por otro mes (ej. 2026-09), no aparece
        res_sep = client.get("/api/dashboard/reconciliation/pending-invoices?month=2026-09&account_rfc=DEGF851127TK1")
        assert not any(i["uuid"] == uuid_tukipa for i in res_sep.json())

        # 4. Vincular el primer pago ($2,770.00)
        match1 = client.post("/api/dashboard/reconciliation/match-manual", data={
            "transaction_id": db_tx1.id,
            "uuid_cfdi": uuid_tukipa
        })
        assert match1.status_code == 200

        # Validar alineación de flujo: tx1 debe haber cambiado a INGRESO
        db.refresh(db_tx1)
        assert db_tx1.tipo == "INGRESO"
        assert db_tx1.status_conciliacion == "CONCILIADO"
        assert db_tx1.uuid_cfdi == uuid_tukipa

        # 5. Consultar /pending-invoices sin include_reconciled: ya no debe salir porque conciliado=True
        res_norec = client.get("/api/dashboard/reconciliation/pending-invoices?month=2026-08&account_rfc=DEGF851127TK1")
        assert not any(i["uuid"] == uuid_tukipa for i in res_norec.json())

        # Consultar /pending-invoices CON include_reconciled=True: debe salir mostrando el monto vinculado y saldo restante
        res_rec = client.get("/api/dashboard/reconciliation/pending-invoices?month=2026-08&include_reconciled=true&account_rfc=DEGF851127TK1")
        assert res_rec.status_code == 200
        target_rec = next((i for i in res_rec.json() if i["uuid"] == uuid_tukipa), None)
        assert target_rec is not None
        assert target_rec["pagos_vinculados_count"] == 1
        assert target_rec["monto_vinculado"] == 2770.00
        assert target_rec["saldo_pendiente"] == 443.20

        # 6. Vincular el segundo pago ($443.20) a la misma factura
        match2 = client.post("/api/dashboard/reconciliation/match-manual", data={
            "transaction_id": db_tx2.id,
            "uuid_cfdi": uuid_tukipa
        })
        assert match2.status_code == 200

        db.refresh(db_tx2)
        assert db_tx2.tipo == "INGRESO"
        assert db_tx2.status_conciliacion == "CONCILIADO"
        assert db_tx2.uuid_cfdi == uuid_tukipa

        # Verificar saldo pendiente actualizado a 0.0
        res_final = client.get("/api/dashboard/reconciliation/pending-invoices?month=2026-08&include_reconciled=true&account_rfc=DEGF851127TK1")
        target_final = next(i for i in res_final.json() if i["uuid"] == uuid_tukipa)
        assert target_final["pagos_vinculados_count"] == 2
        assert target_final["monto_vinculado"] == 3213.20
        assert target_final["saldo_pendiente"] == 0.0
    finally:
        db.close()






